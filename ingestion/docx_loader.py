"""
DOCX Loader Module
------------------
Loads Microsoft Word documents and converts them into a Document object.
"""

from pathlib import Path
from docx import Document as DocxDocument

from ingestion.models import Document


class DOCXLoader:
    """
    Loads DOCX files and extracts text.
    """

    def load(self, file_path: str) -> Document:

        doc = DocxDocument(file_path)

        full_text = ""

        for paragraph in doc.paragraphs:
            full_text += paragraph.text + "\n"

        return Document(
            text=full_text,
            source=file_path,
            file_name=Path(file_path).name,
            file_type="docx",
            page_count=1,
            ocr_used=False,
            metadata={}
        )